"""Constrói o site (docs/), o guia de leitura em PDF e a planilha de acompanhamento a partir de data/catalogo.json.

Uso:  python3 ferramentas/construir.py            (endereço público padrão)
      SITE_URL=https://exemplo.org/bf/ python3 ferramentas/construir.py
Opções: --sem-pdf (pula o guia em PDF).
Requisitos: Python 3.10+, weasyprint (PDF) e openpyxl (planilha). O site é HTML/CSS/JS puro, com links relativos.
"""
import csv, hashlib, html, json, os, re, sys, unicodedata
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, 'docs')
SITE = os.environ.get('SITE_URL', 'https://oaleffemanuel.github.io/biblioth-que-fran-aise/').rstrip('/') + '/'
REPO = 'https://github.com/oaleffemanuel/biblioth-que-fran-aise'
PDF_NAME = 'guia-de-leitura-bibliotheque-francaise.pdf'
XLSX_NAME = 'diario-de-leitura-bibliotheque-francaise.xlsx'
TODAY = date.today().isoformat()
VER = hashlib.sha1(b''.join(open(os.path.join(DOCS, 'assets', f), 'rb').read() for f in ('style.css', 'app.js'))).hexdigest()[:8]

W = json.load(open(os.path.join(ROOT, 'data', 'catalogo.json'), encoding='utf-8'))['obras']
e = lambda s: html.escape(str(s or ''), quote=True)
TITLE_FIX = {'Œuvres complètes - Volume 1': 'Poèmes saturniens, Fêtes galantes, La Bonne Chanson (Œuvres, vol. 1)'}
def clean_title(t):
    t = TITLE_FIX.get(t, t)
    return re.split(r':\s', t)[0].rstrip(' ,;') if ': ' in t else t.rstrip(' ,;')
for i, w in enumerate(W):
    w['i'] = i; w['t'] = clean_title(w['titulo'])

def forma(w):
    g, p = (w['genero'] or ''), w['parcours']
    if w['tipo'] in ('História ilustrada', 'História graduada', 'Leitor para estudantes'): return w['tipo']
    if 'espiritualidade' in p and not g.startswith(('Teatro', 'Poesia')): return 'Espiritualidade'
    if re.search(r'Teatro|Comédia|Tragédia', g): return 'Teatro'
    if re.search(r'Poesia|Fábulas', g): return 'Poesia'
    if re.search(r'Conto|Novela', g): return 'Contos e novelas'
    if re.search(r'Filosofia|Hist|Memórias|Máximas|Moralistas|Biografia|Ensaio', g): return 'Ensaio, filosofia e história'
    return 'Romance'
for w in W: w['forma'] = forma(w)

# ───────────────────────── o percurso: 6 étapes ─────────────────────────
ETAPES = {
    1: dict(nome='Premiers pas', pt='Primeiros passos', nivel='A1',
            curta='Livrinhos ilustrados de poucas frases e as primeiras histórias A1, com narração na página oficial.',
            perfil='Você está começando: conhece saudações, números e o presente de verbos como être, avoir e aller.',
            objetivo='Ler as primeiras histórias inteiras e se acostumar com o som do francês escrito.',
            como='Leia cada livrinho duas vezes. Nas histórias MeloLingua, ouça a narração na página oficial depois de ler. Anote no máximo cinco palavras por história.',
            meta='As 20 leituras em 2 a 3 semanas.', avancar='Você entende uma história A1 nova na primeira leitura, sem dicionário.'),
    2: dict(nome='Petites histoires', pt='Pequenas histórias', nivel='A1–A2',
            curta='Histórias mais longas, fábulas africanas e os primeiros leitores escolares, como ABC: Petits Contes (com audiolivro).',
            perfil='Você reconhece o passé composé e acompanha diálogos simples.',
            objetivo='Ler por prazer e em quantidade: é aqui que a leitura extensiva começa de verdade.',
            como='Uma leitura por dia. Não pare a cada palavra: se entendeu a ideia, siga em frente. Ao terminar, escreva o résumé en trois phrases.',
            meta='Pelo menos 12 das 18 leituras em 3 a 4 semanas.', avancar='Você lê uma história A2 nova com áudio e entende quase tudo.'),
    3: dict(nome='Le pont', pt='A ponte', nivel='A2–B1',
            curta='Histórias B1 e leitores escolares do início do século XX: contos de fadas, lendas, uma volta pela França e a história do país.',
            perfil='Você lê histórias A2 com facilidade e quer textos com mais enredo.',
            objetivo='Ganhar fôlego para textos de 20 a 60 minutos e se acostumar com o imparfait na narração.',
            como='Divida os leitores escolares em capítulos de 15 minutos. Sublinhe só palavras que se repetem: são as que valem a pena aprender.',
            meta='Uma história B1 por semana e um leitor escolar em andamento.', avancar='Você termina um capítulo de 20 minutos sem perder o fio da história.'),
    4: dict(nome='Contes et nouvelles', pt='Contos e novelas', nivel='B1–B2',
            curta='Os grandes contistas: Maupassant, Daudet, Mérimée, Flaubert e os contos filosóficos de Voltaire. Vários com audiolivro.',
            perfil='Você lê B1 com conforto e quer o primeiro contato com a literatura original.',
            objetivo='Ler os primeiros textos originais, curtos o bastante para terminar, e conhecer o passé simple.',
            como='Escolha contos com audiolivro (Lettres de mon moulin, Contes de la bécasse, Boule de suif): leia um conto, depois ouça. Aceite não entender tudo.',
            meta='Uma coletânea de contos inteira em 1 a 2 meses.', avancar='Você reconhece as formas do passé simple (il fut, il eut, il dit, il fit) sem parar.'),
    5: dict(nome='Le roman', pt='O romance', nivel='B2',
            curta='Romances do século XIX: Verne, Dumas, Hugo, Balzac, Zola, Molière no teatro e 100 romances do corpus ELTeC para baixar.',
            perfil='Você já leu contos originais inteiros e quer um livro longo.',
            objetivo='Ler o primeiro romance inteiro em francês e criar o hábito da leitura longa.',
            como='Comece por um romance curto ou com audiolivro (Tartarin de Tarascon, Voyage au centre de la Terre). Leia um capítulo por dia; alterne com uma peça de Molière.',
            meta='Um romance a cada um ou dois meses.', avancar='A sintaxe do século XIX não assusta mais e você quer encarar os clássicos.'),
    6: dict(nome='Les classiques', pt='Os clássicos', nivel='C1',
            curta='Racine, Corneille, Pascal, os moralistas, a poesia (Baudelaire, Hugo, Rimbaud), a história e os grandes textos espirituais.',
            perfil='Você lê romances do século XIX com conforto.',
            objetivo='Conhecer as obras que formaram a língua e o pensamento franceses.',
            como='Leia aos poucos e com apoio: uma cena por dia de Phèdre, alguns fragmentos dos Pensées, um poema por vez. Use uma tradução em português ao lado quando quiser.',
            meta='Um clássico por vez, sem pressa.', avancar='Não há próxima etapa. Bonne lecture !'),
}
TRILHAS = {'espiritualidade': ("L'âme", 'Espiritualidade católica', 'Da Imitação de Cristo à História de uma alma, passando por Francisco de Sales, Santo Agostinho, Bossuet, Pascal, Péguy e Claudel.'),
           'pensamento': ('La pensée', 'Filosofia e história', 'Voltaire, Montesquieu, Descartes, os moralistas, Rousseau, Tocqueville, Michelet e Chateaubriand.'),
           'poesia-teatro': ('La scène et le vers', 'Teatro e poesia', 'Molière, Racine e Corneille; La Fontaine, Hugo, Baudelaire, Verlaine, Rimbaud e Apollinaire.')}
for t in ETAPES: ETAPES[t]['obras'] = [w for w in W if w['etapa'] == t]
INKLINGO = 'https://www.inklingo.app/french/stories/'
EXTERNOS = {1: [('Inklingo — histórias A0 e A1', INKLINGO + 'a0'), ('MeloLingua — todas as histórias A1, com áudio', 'https://melolingua.com/learn-french')],
            2: [('Inklingo — histórias A2', INKLINGO + 'a2')], 3: [('Inklingo — histórias B1', INKLINGO + 'b1')], 4: [('Inklingo — histórias B2', INKLINGO + 'b2')]}

def find(title, author=None):
    for w in W:
        if (w['t'] == title or w['titulo'] == title) and (author is None or w['autor'] == author): return w
    sys.exit(f'Obra não encontrada no catálogo: {title}')

def picks(t):
    if t == 5:
        short = [w for w in ETAPES[5]['obras'] if w['fonte'] == 'ELTeC-fra'][:5]
        return [('Para começar: romances curtos ou com audiolivro', [find(x) for x in ['Tartarin de Tarascon', "Le Dernier Jour d'un Condamné", 'Voyage au centre de la Terre', 'Eugénie Grandet', 'Cinq semaines en ballon']]),
                ('Teatro de Molière (peças curtas)', [find(x) for x in ['Les Précieuses ridicules', 'Le médecin malgré lui', "L'Avare", 'Le Bourgeois gentilhomme', 'Tartuffe ou l’Imposteur']]),
                ('Os grandes romances', [find(x) for x in ['Les trois mousquetaires', 'Le comte de Monte-Cristo', 'Les Misérables', 'Germinal', 'Le rouge et le noir']]),
                ('Corpus ELTeC: os mais curtos dos 100 romances (PDF aqui)', short)]
    if t == 6:
        return [('Uma ordem possível', [find(x) for x in ['Phèdre', 'Le Cid', 'La Chanson de Roland', 'Réflexions ou sentences et maximes morales', 'Lettres persanes',
                                                           'Discours de la méthode', 'Les Fleurs du Mal', 'Les Contemplations', 'Pensées', "Mémoires d'Outre-Tombe"]])]
    return [('Leituras desta etapa', ETAPES[t]['obras'])]

def dur(m):
    if not m: return None
    if m < 60: return f'~{m} min'
    h = m / 60
    return f'~{round(h)} h' if h >= 2 else f'~{h:.1f} h'.replace('.', ',')
def autor_line(w):
    if w['tradutor'] and w['fonte'] == 'Wikisource (fr)' and w['autor'] == 'Anônimo': return 'Anônimo · tradução de 1911'
    return w['autor'] or w['fonte']
def nivel_label(w):
    if w['cefr_oficial']: return f"{w['cefr_oficial']} oficial"
    return ETAPES[w['etapa']]['nivel']   # estimativa coerente com a etapa

def links(w):
    """Mesma lógica de botões da Biblioteca Italiana: Baixar PDF / Baixar EPUB / Acessar; áudio à parte."""
    out = []
    if w['arquivo']: out.append(('Baixar PDF', w['arquivo'], True))
    ep = w.get('epub_urls') or []
    if len(ep) == 1: out.append(('Baixar EPUB', ep[0], False))
    if w['url_fonte'] and w['fonte'] != 'ELTeC-fra':
        au = w['audio'] or {}
        lab = 'Ler e ouvir' if au.get('fonte') == 'MeloLingua' else ('Acessar' + (f' ({len(ep)} volumes)' if len(ep) > 1 else ''))
        out.append((lab, w['url_fonte'], False))
    elif not w['arquivo']: out.append(('Acessar', w['url_fonte'], False))
    return out
def audio_link(w):
    a = w['audio'] or {}
    return (a['url'], a.get('duracao')) if a.get('fonte') == 'LibriVox' else None

ICON_AUDIO = '<svg aria-hidden="true" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round"><path d="M3 14v-2a9 9 0 0 1 18 0v2"/><rect x="3" y="14" width="4" height="7" rx="1.5"/><rect x="17" y="14" width="4" height="7" rx="1.5"/></svg>'
ICON_DL = '<svg aria-hidden="true" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round"><path d="M12 3v12m0 0-5-5m5 5 5-5M4 21h16"/></svg>'

# ───────────────────────── layout ─────────────────────────
NAV = [('./', 'Início'), ('catalogo.html', 'Catálogo'), ('plano.html', 'Percurso de leitura'), ('guia.html', 'Como usar'), ('sobre.html', 'Sobre e direitos')]
def page(fname, title, desc, body, extra_head=''):
    nav = ''.join(f'<a href="{h}"{" aria-current=\"page\"" if h == fname or (h == "./" and fname == "index.html") else ""}>{t}</a>' for h, t in NAV)
    canon = SITE + ('' if fname == 'index.html' else fname)
    full = title if fname == 'index.html' else f'{title} · Bibliothèque Française'
    seo = ('<base href="' + re.sub(r'^https?://[^/]+', '', SITE) + '">\n<meta name="robots" content="noindex">') if fname == '404.html' else \
          f'<link rel="canonical" href="{canon}">\n<meta property="og:url" content="{canon}">'
    return f'''<!doctype html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(full)}</title>
<meta name="description" content="{e(desc)}">
{seo}
<meta property="og:type" content="website">
<meta property="og:site_name" content="Bibliothèque Française">
<meta property="og:title" content="{e(full)}">
<meta property="og:description" content="{e(desc)}">
<meta property="og:image" content="{SITE}assets/og.png">
<meta property="og:locale" content="pt_BR">
<meta name="twitter:card" content="summary_large_image">
<meta name="theme-color" content="#14213d">
<link rel="icon" href="assets/favicon.svg" type="image/svg+xml">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Cormorant+Garamond:ital,wght@0,500;0,600;1,500;1,600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="assets/style.css?v={VER}">
{extra_head}</head>
<body>
<a class="skip" href="#conteudo">Pular para o conteúdo</a>
<header class="site-header">
<div class="eco"><a href="https://lg.korium.co/">← Todas as bibliotecas</a></div>
<div class="wrap">
<a class="brand" href="./"><img src="assets/favicon.svg" alt="" width="34" height="34"><span>Bibliothèque Française<small>lire pour apprendre</small></span></a>
<button class="nav-toggle" type="button" id="nav-toggle" aria-expanded="false" aria-controls="menu" hidden>Menu</button>
<nav class="nav" id="menu" aria-label="Principal">{nav}</nav>
</div></header>
<main id="conteudo">
{body}
</main>
<footer class="site-footer"><div class="wrap">
<div><p><b>Bibliothèque Française</b> · leituras em francês, do A1 aos clássicos.</p>
<p>Projeto pessoal, gratuito e sem fins lucrativos, feito por um brasileiro que estuda francês. Não é afiliado às fontes citadas.</p>
<p>Da mesma família: <a href="https://oaleffemanuel.github.io/biblioteca-italiana/">Biblioteca Italiana</a>.</p></div>
<div><p><a href="{REPO}">Código e dados no GitHub</a> · <a href="{REPO}/issues">Avisar um problema</a></p>
<p>Textos do projeto: CC BY 4.0 · Atualizado em {TODAY}</p></div>
</div></footer>
<script src="assets/app.js?v={VER}" defer></script>
</body>
</html>
'''

def card(w):
    t = w['etapa']; lvl = nivel_label(w)
    tags = [f'<span class="tag t">Etapa {t}{" · " + e(lvl) if lvl else ""}</span>']
    if w['nivel_fonte']: tags.append(f'<span class="tag">{e(w["nivel_fonte"].split(" — ")[0])} (ASb)</span>')
    if w['audio']: tags.append(f'<span class="tag audio">{ICON_AUDIO} {"Audiolivro" if w["audio"]["fonte"] == "LibriVox" else "Áudio"}</span>')
    if w['arquivo']: tags.append(f'<span class="tag dl-tag">{ICON_DL} PDF aqui</span>')
    acts = []
    for k, (lab, href, dl) in enumerate(links(w)):
        cls = 'btn small accent' if dl else ('btn small' if k == 0 else 'btn small ghost')
        acts.append(f'<a class="{cls}" href="{e(href)}"' + (' download' if dl else ' target="_blank" rel="noopener"') + f'>{e(lab)}</a>')
    al = audio_link(w)
    if al: acts.append(f'<a class="btn small ghost" href="{e(al[0])}" target="_blank" rel="noopener">Audiolivro</a>')
    acts.append('<button type="button" class="mark want" data-kind="quero" aria-pressed="false">Quero ler</button>')
    acts.append('<button type="button" class="mark" data-kind="lido" aria-pressed="false">✓ Lido</button>')
    meta = ' · '.join(x for x in [w['forma'], dur(w['minutos']), f"{w['paginas']} págs." if w['paginas'] else None, w['fonte']] if x)
    q = ' '.join(str(x) for x in [w['t'], w['autor'], w['tradutor'], w['fonte'], w['forma'], w['genero'], lvl, w['resumo']] if x)
    qn = ''.join(c for c in unicodedata.normalize('NFD', q.lower()) if unicodedata.category(c) != 'Mn')
    resumo = f'<p class="meta"><i>{e(w["resumo"])}</i></p>' if w['resumo'] else ''
    return (f'<article class="obra c{t}" id="{e(w["id"])}" data-i="{w["i"]}" data-etapa="{t}" data-forma="{e(w["forma"])}" '
            f'data-fonte="{e(w["fonte"])}" data-audio="{int(bool(w["audio"]))}" data-pdf="{int(bool(w["arquivo"]))}" data-min="{w["minutos"] or 99999}" '
            f'data-trilha="{" ".join(w["parcours"])}" data-t="{e(w["t"])}" data-q="{e(qn)}">'
            f'<div class="top">{"".join(tags)}</div><h3>{e(w["t"])}</h3><p class="autor">{e(autor_line(w))}</p>'
            f'<p class="meta">{e(meta)}</p>{resumo}<div class="acts">{"".join(acts)}</div></article>')

def etape_cards():
    out = []
    for t, d in ETAPES.items():
        n = len(d['obras']); a = sum(bool(w['audio']) for w in d['obras'])
        out.append(f'<a class="tappa c{t}" href="catalogo.html?etapa={t}"><span class="num">Étape {t}</span><h3>{e(d["nome"])}</h3>'
                   f'<span class="lvl">{e(d["pt"])} · {e(d["nivel"])}</span><p>{e(d["curta"])}</p>'
                   f'<div class="count">{n} obras{f" · {a} com áudio" if a else ""}</div></a>')
    return '<div class="tappe">' + ''.join(out) + '</div>'

N = len(W); NA = sum(bool(w['audio']) for w in W); NP = sum(bool(w['arquivo']) for w in W)

def p_index():
    ld = json.dumps({'@context': 'https://schema.org', '@type': 'WebSite', 'name': 'Bibliothèque Française', 'url': SITE, 'inLanguage': 'pt-BR',
                     'description': 'Biblioteca gratuita de leituras em francês organizadas por nível, do A1 aos clássicos.'}, ensure_ascii=False)
    tr = ''.join(f'<li><a href="catalogo.html?trilha={k}"><b class="it">{e(a)}</b> — {e(b)}</a><br><small>{e(c)}</small></li>' for k, (a, b, c) in TRILHAS.items())
    body = f'''<section class="hero"><div class="wrap">
<p class="kicker">Bienvenue ! Une bibliothèque pour apprendre le français en lisant.</p>
<h1>Leia francês no seu nível.</h1>
<p class="lead">{N} leituras gratuitas, de livrinhos ilustrados para iniciantes até Racine, Pascal e Hugo, reunidas de fontes confiáveis e organizadas em seis etapas. Sem cadastro e sem precisar saber nada de GitHub.</p>
<ul class="stats"><li><b>{N}</b>obras</li><li><b>{NA}</b>com áudio</li><li><b>{NP}</b>PDFs para baixar aqui</li><li><b>6</b>etapas de leitura</li></ul>
<div class="cta"><a class="btn accent" href="catalogo.html">Explorar o catálogo</a><a class="btn ghost" href="plano.html#onde-comecar">Não sei meu nível</a></div>
</div></section>
<section class="wrap"><h2>O percurso, em seis <i>étapes</i></h2>
<p>Escolha a etapa que combina com o seu nível de hoje; cada cartão abre o catálogo já filtrado.</p>
{etape_cards()}
</section>
<section class="wrap"><h2>Três trilhas temáticas</h2>
<p>Para quem já está nas etapas 5 e 6 e quer ler por interesse:</p>
<ul class="steps trilhas">{tr}</ul>
</section>
<section class="wrap"><h2>Como funciona</h2>
<ol class="steps">
<li><b>Escolha uma leitura.</b><br>Filtre por etapa, por áudio ou por duração. Ou deixe o catálogo sortear uma.</li>
<li><b>Baixe ou acesse.</b><br>"Baixar PDF" baixa o livro deste site; "Acessar" abre o livro na fonte oficial.</li>
<li><b>Leia e ouça.</b><br>Onde houver áudio, leia primeiro e depois ouça. Anote até 5 palavras e escreva três frases sobre o texto.</li>
<li><b>Marque seu progresso.</b><br>Use "Quero ler" e "Lido" no site ou a planilha de acompanhamento.</li>
</ol></section>
<section class="wrap"><h2>Para baixar</h2>
<div class="downloads">
<div class="dl"><h3>Guia de leitura (PDF)</h3><p>O percurso completo: como começar, como escolher livros, ler ouvindo, vocabulário, resumos e o caminho até a literatura original, com as obras da biblioteca.</p><a class="btn accent" href="telechargements/{PDF_NAME}" download>{ICON_DL} Baixar o guia</a> <a class="btn ghost" href="plano.html">Ler online</a></div>
<div class="dl"><h3>Diário de leitura (planilha)</h3><p>As {N} obras com links, status, datas e notas, mais painel de progresso, sessões, vocabulário com revisão e resumos.</p><a class="btn accent" href="telechargements/{XLSX_NAME}" download>{ICON_DL} Baixar a planilha</a> <a class="btn ghost" href="guia.html#planilha">Como usar</a></div>
</div>
<p class="note">Só hospedamos arquivos cuja licença permite redistribuição; o resto leva à fonte oficial. <a href="sobre.html#direitos">Entenda por quê.</a></p>
</section>'''
    return page('index.html', 'Bibliothèque Française — leituras em francês por nível, do A1 aos clássicos',
                f'{N} leituras gratuitas em francês organizadas por nível: livrinhos ilustrados, histórias A1–B2 com áudio, contos, romances e clássicos. Guia de leitura em PDF e planilha de acompanhamento.',
                body, f'<script type="application/ld+json">{ld}</script>\n')

def p_catalogo():
    fontes = sorted({w['fonte'] for w in W})
    formas = ['História ilustrada', 'História graduada', 'Leitor para estudantes', 'Contos e novelas', 'Romance', 'Teatro', 'Poesia', 'Ensaio, filosofia e história', 'Espiritualidade']
    chips_t = ''.join(f'<button type="button" class="chip-btn c{t}" data-group="etapa" data-value="{t}" aria-pressed="false">{t} · {e(d["nome"])}</button>' for t, d in ETAPES.items())
    chips_f = ''.join(f'<button type="button" class="chip-btn" data-group="forma" data-value="{e(f)}" aria-pressed="false">{e(f)}</button>' for f in formas)
    tr_opts = ''.join(f'<option value="{k}">{e(b)}</option>' for k, (a, b, c) in TRILHAS.items())
    body = f'''<div class="wrap"><h1>Catálogo</h1>
<p>{N} obras em ordem sugerida de leitura, da mais fácil à mais difícil. Combine os filtros à vontade; o endereço da página guarda a busca para você compartilhar.</p></div>
<div class="filters"><div class="wrap">
<div class="search">
<input id="q" type="search" placeholder="Buscar título, autor ou tema (ex.: Maupassant, Molière, chat)" aria-label="Buscar" autocomplete="off">
<select id="fonte" aria-label="Fonte"><option value="">Todas as fontes</option>{"".join(f'<option>{e(f)}</option>' for f in fontes)}</select>
<select id="trilha" aria-label="Trilha"><option value="">Todas as trilhas</option>{tr_opts}</select>
<select id="ordem" aria-label="Ordenar"><option value="sugerida">Ordem sugerida</option><option value="curtas">Mais curtas primeiro</option><option value="longas">Mais longas primeiro</option><option value="titulo">Título (A–Z)</option></select>
</div>
<button type="button" class="filters-toggle" id="filtros-botao" aria-expanded="true" aria-controls="filtros-mais" hidden>Mais filtros</button>
<div class="filters-more" id="filtros-mais">
<div class="chips" role="group" aria-label="Etapa">{chips_t}</div>
<div class="chips" role="group" aria-label="Forma">{chips_f}</div>
<div class="toggles">
<label><input type="checkbox" id="f-audio"> Com áudio</label>
<label><input type="checkbox" id="f-pdf"> PDF para baixar aqui</label>
<label>Minha lista: <select id="lista"><option value="">todas</option><option value="quero">quero ler</option><option value="lido">já li</option><option value="nao-lido">ainda não li</option></select></label>
</div></div></div></div>
<div class="wrap">
<div class="result-bar"><span id="count" aria-live="polite">{N} obras</span>
<span><button type="button" class="btn small ghost" id="sorteio">🎲 Sortear uma leitura</button> <button type="button" class="btn small ghost" id="limpar">Limpar filtros</button></span></div>
<div class="grid" id="grid">{"".join(card(w) for w in W)}</div>
<p class="empty" id="empty" hidden>Nenhuma obra com esses filtros. Tente limpar algum deles.</p>
<p class="note">"Quero ler" e "Lido" ficam salvos só neste navegador, sem cadastro. O tempo é uma estimativa para quem está aprendendo. "A1 oficial" indica nível declarado pela fonte (MeloLingua); os demais níveis são estimativas nossas. Os níveis do African Storybook (ASb) são da própria fonte e não são CEFR. Catálogo em <a href="donnees/catalogo.csv">CSV</a> e <a href="donnees/catalogo.json">JSON</a>.</p>
</div>'''
    return page('catalogo.html', 'Catálogo de leituras em francês por nível',
                f'Catálogo com {N} leituras em francês: filtre por nível, áudio, duração e forma. Histórias para iniciantes, contos, romances, teatro, poesia e espiritualidade.', body)

def plan_body(pdf=False):
    link = lambda w: f'<a href="{e((SITE + w["arquivo"]) if (pdf and w["arquivo"]) else (w["arquivo"] or w["url_fonte"]))}">{e(w["t"])}</a>'
    def acesso(w):
        if w['arquivo']: return 'PDF no site'
        return 'EPUB no Gutenberg' if w.get('epub_urls') else e(w['fonte'].split(' +')[0])
    def table(ws):
        rows = ''.join(f'<tr><td>{link(w)}</td><td>{e(autor_line(w))}</td><td>{e(nivel_label(w) or "—")}</td><td>{dur(w["minutos"]) or "—"}</td>'
                       f'<td>{"audiolivro" if audio_link(w) else ("áudio na fonte" if w["audio"] else "—")}</td><td>{acesso(w)}</td></tr>' for w in ws)
        return f'<div class="table-scroll"><table><thead><tr><th>Obra</th><th>Autor</th><th>Nível</th><th>Tempo</th><th>Áudio</th><th>Onde</th></tr></thead><tbody>{rows}</tbody></table></div>'
    sec = ''
    for t, d in ETAPES.items():
        groups = ''.join(f'<h3>{e(h)}</h3>{table(ws)}' for h, ws in picks(t))
        ext = ''.join(f'<li><a href="{u}">{e(n)}</a></li>' for n, u in EXTERNOS.get(t, []))
        ext = f'<p class="ext"><b>Para ler mais (sites externos, só online):</b></p><ul class="ext">{ext}</ul>' if ext else ''
        more = f'<p>Todas as {len(d["obras"])} obras desta etapa: {SITE}catalogo.html?etapa={t}</p>' if pdf else f'<p><a href="catalogo.html?etapa={t}">Ver as {len(d["obras"])} obras desta etapa no catálogo →</a></p>'
        sec += f'''<section class="tappa-sec c{t}" id="etapa-{t}">
<div class="tappa-head"><h2>Étape {t} · <i>{e(d["nome"])}</i></h2><span class="tag t">{e(d["pt"])} · {e(d["nivel"])}</span></div>
<p>{e(d["curta"])}</p>
<div class="two"><div class="box"><b>Para quem</b><br>{e(d["perfil"])}<br><br><b>Objetivo</b><br>{e(d["objetivo"])}</div>
<div class="box"><b>Como ler</b><br>{e(d["como"])}<br><br><b>Meta</b><br>{e(d["meta"])}</div></div>
{groups}{ext}
<p class="note"><b>Sinal para avançar:</b> {e(d["avancar"])}</p>{more}
</section>'''
    trilhas = ''
    for k, (a, b, c) in TRILHAS.items():
        ws = [w for w in W if k in w['parcours']]
        trilhas += f'<h3><i>{e(a)}</i> — {e(b)}</h3><p>{e(c)}</p>' + table(sorted(ws, key=lambda w: (w['etapa'], w['minutos'] or 9e9))[:12])
    return f'''
<p class="lead">Um caminho prático para sair do A1 e chegar à literatura francesa original, usando as obras desta biblioteca. Com 20 a 30 minutos por dia, as etapas 1 a 4 levam de 4 a 6 meses.</p>

<h2 id="onde-comecar">Onde começar?</h2>
<p>Leia o começo de cada texto abaixo, em ordem, e pare no primeiro em que você precisar do dicionário mais de uma vez por frase.</p>
<ol>
<li>{link(find("Le porridge"))} (livrinho ilustrado). Se foi difícil, comece pela <b>Étape 1</b>.</li>
<li>{link(find("Une rencontre au café"))} (A1 oficial). Se foi tranquilo, vá para a <b>Étape 2</b>.</li>
<li>{link(find("La rencontre au marché"))} (B1 oficial). Se foi tranquilo, comece pela <b>Étape 3</b>.</li>
<li>{link(find("Micromégas"))} (Voltaire, texto original). Se foi tranquilo, vá para a <b>Étape 4</b>, ou direto à 5 se já lê com facilidade.</li>
</ol>
<p>Na dúvida, comece uma etapa abaixo. Ler fácil demais custa pouco; ler difícil demais faz desistir.</p>

<h2 id="metodo">O método: quatro hábitos</h2>
<div class="two">
<div class="box"><b>1. Leia fácil e leia muito</b><br>Quase todo o seu tempo deve ir para textos que você entende quase inteiros. É a <i>leitura extensiva</i>: o vocabulário se fixa pela repetição. <b>Regra prática:</b> se uma página tem mais de 5 a 8 palavras desconhecidas que atrapalham a história, o texto está difícil demais — volte uma etapa.</div>
<div class="box"><b>2. Leia e ouça</b><br>Para os textos com áudio, faça três passagens: <b>(1)</b> leia sem áudio; <b>(2)</b> leia ouvindo; <b>(3)</b> só ouça, em outro momento do dia. O francês escrito e o falado são muito diferentes; o áudio liga os dois.</div>
<div class="box"><b>3. Poucas palavras, bem escolhidas</b><br>No máximo <b>cinco palavras por texto</b>, sempre com a frase em que apareceram. Revise depois de 1, 3 e 7 dias (a planilha calcula as datas). Prefira as palavras que se repetem.</div>
<div class="box"><b>4. Escreva o <i>résumé en trois phrases</i></b><br>Ao terminar, três frases em francês: <b>Qui ?</b> (quem e onde) · <b>Que se passe-t-il ?</b> (o que acontece) · <b>Comment ça finit ?</b> (o fim, ou o que você achou). Exemplo inventado: <i>Marie perd son chat dans le quartier. Elle le cherche toute la journée avec son voisin. Le soir, le chat dort dans sa cuisine.</i></div>
</div>

<h3>Uma semana típica (30 minutos por dia)</h3>
<div class="table-scroll keep"><table><thead><tr><th>Dia</th><th>O que fazer</th></tr></thead><tbody>
<tr><td>Segunda a sexta</td><td>20 min de leitura nova + 10 min ouvindo um texto já lido, ou revisando palavras</td></tr>
<tr><td>Sábado</td><td>Escrever os <i>résumés</i> da semana e revisar o vocabulário</td></tr>
<tr><td>Domingo</td><td>Livre: releia o texto favorito da semana, só ouvindo</td></tr>
</tbody></table></div>

<h3>Ferramentas</h3>
<p>A <b>planilha de acompanhamento</b> ({SITE}telechargements/{XLSX_NAME}) traz todas as obras, com status, datas, nota, sessões, vocabulário com revisão e espaço para os <i>résumés</i>. O <b>catálogo online</b> ({SITE}catalogo.html) permite filtrar e sortear leituras.</p>

<h2 id="etapas">As seis etapas</h2>
<p>Só as histórias MeloLingua têm nível CEFR <b>oficial</b> (A1–B2). Os níveis do African Storybook são da própria fonte (de "primeiras palavras" a "parágrafos mais longos") e não equivalem ao CEFR. Todos os outros níveis deste guia são estimativas nossas.</p>
{sec}

<h2 id="trilhas">Três trilhas temáticas</h2>
<p>Depois da Étape 4, você pode seguir também por interesse. As listas abaixo mostram as primeiras obras de cada trilha, da mais acessível à mais difícil.</p>
{trilhas}

<h2 id="dicas">Dicas para brasileiros</h2>
<h3>Falsos amigos (<i>faux amis</i>) frequentes</h3>
<div class="table-scroll keep"><table><thead><tr><th>Francês</th><th>Significa</th><th>Não é</th></tr></thead><tbody>
<tr><td>attendre</td><td>esperar</td><td>atender (répondre, servir)</td></tr>
<tr><td>la journée</td><td>o dia (duração)</td><td>jornada de trabalho</td></tr>
<tr><td>le parent</td><td>pai ou mãe; parente</td><td>só "parente"</td></tr>
<tr><td>rester</td><td>ficar, permanecer</td><td>restar (seulement)</td></tr>
<tr><td>la salade</td><td>alface; salada</td><td>—</td></tr>
<tr><td>la serviette</td><td>guardanapo; toalha</td><td>—</td></tr>
<tr><td>sensible</td><td>sensível</td><td>sensato (raisonnable)</td></tr>
<tr><td>prétendre</td><td>afirmar</td><td>pretender (avoir l'intention)</td></tr>
<tr><td>le magasin</td><td>a loja</td><td>a revista (le magazine)</td></tr>
<tr><td>la fac</td><td>a universidade</td><td>a faca (le couteau)</td></tr>
<tr><td>large</td><td>largo, amplo</td><td>—</td></tr>
<tr><td>le collège</td><td>escola dos 11 aos 15 anos</td><td>colégio em geral / faculdade</td></tr>
</tbody></table></div>
<h3>O <i>passé simple</i>: o tempo verbal da literatura</h3>
<p>A partir da Étape 4 quase toda narração está no <i>passé simple</i>. Basta reconhecê-lo:</p>
<div class="table-scroll keep"><table><thead><tr><th>Infinitivo</th><th>il / elle</th><th>ils / elles</th><th>Português</th></tr></thead><tbody>
<tr><td>être</td><td>fut</td><td>furent</td><td>foi / foram</td></tr>
<tr><td>avoir</td><td>eut</td><td>eurent</td><td>teve / tiveram</td></tr>
<tr><td>dire</td><td>dit</td><td>dirent</td><td>disse / disseram</td></tr>
<tr><td>faire</td><td>fit</td><td>firent</td><td>fez / fizeram</td></tr>
<tr><td>venir</td><td>vint</td><td>vinrent</td><td>veio / vieram</td></tr>
<tr><td>voir</td><td>vit</td><td>virent</td><td>viu / viram</td></tr>
<tr><td>prendre</td><td>prit</td><td>prirent</td><td>pegou / pegaram</td></tr>
<tr><td>parler · finir · vendre</td><td>parla · finit · vendit</td><td>parlèrent · finirent · vendirent</td><td>regulares em -a, -it</td></tr>
</tbody></table></div>
<h3>Francês antigo e clássico: o que muda</h3>
<p>Nos textos dos séculos XVII e XVIII (Étape 6), algumas grafias parecem estranhas: <i>-oit, -oient</i> no lugar de <i>-ait, -aient</i> (<i>il étoit</i> = il était), <i>sçavoir</i> = savoir, <i>tems</i> = temps. No teatro clássico o texto é em versos de 12 sílabas (alexandrinos): leia em voz alta, devagar.</p>
'''

def p_plano():
    body = f'''<div class="wrap narrow prose">
<h1>Percurso de leitura</h1>
<p><a class="btn accent" href="telechargements/{PDF_NAME}" download>{ICON_DL} Baixar o guia em PDF</a> <a class="btn ghost" href="telechargements/{XLSX_NAME}" download>{ICON_DL} Planilha de acompanhamento</a></p>
{plan_body(False)}
</div>'''
    return page('plano.html', 'Percurso de leitura em francês, do A1 aos clássicos',
                'Percurso gratuito para aprender francês lendo: seis etapas do A1 ao C1 com obras reais, método de leitura com áudio, trilhas temáticas e dicas para brasileiros. Guia em PDF.', body)

def p_guia():
    body = f'''<div class="wrap narrow prose">
<h1>Como usar</h1>
<p class="lead">Você não precisa de conta, cadastro nem saber o que é GitHub. Tudo funciona no navegador, inclusive no celular.</p>
<h2>Em 3 minutos</h2>
<ol><li>Descubra seu ponto de partida com o <a href="plano.html#onde-comecar">teste rápido</a>.</li>
<li>Abra o <a href="catalogo.html">catálogo</a> e clique na sua etapa.</li>
<li>Escolha uma obra e clique em <b>Baixar PDF</b>, <b>Baixar EPUB</b> ou <b>Acessar</b>.</li>
<li>Ao terminar, clique em <b>✓ Lido</b>. O site lembra para você.</li></ol>
<h2>Os botões de cada obra</h2>
<ul><li><b>Baixar PDF</b>: baixa o livro deste site (só para obras com licença que permite).</li>
<li><b>Baixar EPUB</b>: baixa o livro direto do Project Gutenberg, para leitores de e-book e celulares.</li>
<li><b>Acessar</b>: abre o livro na fonte oficial. <b>Ler e ouvir</b>: abre a história MeloLingua com o áudio.</li>
<li><b>Audiolivro</b>: abre a gravação gratuita no LibriVox.</li></ul>
<h2>O catálogo</h2>
<ul><li><b>Etapas (1 a 6)</b>: do mais fácil ao mais difícil. Dá para marcar várias.</li>
<li><b>Forma</b>: história ilustrada, história graduada, leitor escolar, contos, romance, teatro, poesia, ensaio, espiritualidade.</li>
<li><b>Trilha</b>: espiritualidade, filosofia e história, teatro e poesia.</li>
<li><b>Sortear uma leitura</b>: escolhe uma obra entre as filtradas que você ainda não leu.</li>
<li><b>Minha lista</b>: "Quero ler" e "Lido" ficam salvos <i>neste navegador</i>. Se limpar os dados ou trocar de aparelho, a lista some; para um registro permanente, use a planilha.</li>
<li><b>Compartilhar uma busca</b>: copie o endereço da página (ex.: <a href="catalogo.html?etapa=1&amp;audio=1">etapa 1 com áudio</a>).</li></ul>
<h2>Lendo no celular</h2>
<ul><li><b>PDFs</b> abrem no navegador; para guardar, use "Compartilhar → Salvar em Arquivos" (iPhone) ou "Baixar" (Android).</li>
<li><b>EPUB</b> abre no Apple Books, no Google Play Livros ou em leitores como Kobo.</li>
<li>Os <b>livrinhos do African Storybook</b> também podem ser lidos online, com as ilustrações, em "Acessar".</li></ul>
<h2 id="planilha">A planilha de acompanhamento</h2>
<p><a class="btn accent" href="telechargements/{XLSX_NAME}" download>{ICON_DL} Baixar a planilha</a></p>
<p>Abre no Excel, no LibreOffice ou no <b>Google Planilhas</b> (<i>Arquivo → Importar → Fazer upload</i>). Seis abas: <b>Comece aqui</b>, <b>Painel</b> (progresso por etapa, calculado sozinho), <b>Catálogo</b> (as {N} obras com link e status), <b>Sessões</b>, <b>Vocabulário</b> (com datas de revisão) e <b>Résumés</b>.</p>
<h2>Perguntas frequentes</h2>
<details><summary>É gratuito?</summary><p>Sim: as obras, o site, o guia e a planilha.</p></details>
<details><summary>Por que algumas obras só têm "Acessar"?</summary><p>Porque acesso gratuito não é permissão para redistribuir. Só hospedamos arquivos cuja licença permite; para o resto, levamos você à fonte oficial. Veja <a href="sobre.html#direitos">a política de direitos</a>.</p></details>
<details><summary>O que significam A1, A2, B1, B2, C1?</summary><p>São os níveis do Quadro Europeu Comum de Referência. Aqui, só as histórias MeloLingua têm nível oficial; os demais são estimativas nossas.</p></details>
<details><summary>Encontrei um link quebrado ou um erro.</summary><p>Avise por uma <a href="{REPO}/issues">issue no GitHub</a>.</p></details>
</div>'''
    return page('guia.html', 'Como usar', 'Guia rápido da Bibliothèque Française: botões, catálogo, leitura no celular, planilha de acompanhamento e perguntas frequentes.', body)

def p_sobre():
    cnt = {}
    for w in W: cnt[w['fonte']] = cnt.get(w['fonte'], 0) + 1
    rows = [('African Storybook', 'https://www.africanstorybook.org', 'Livrinhos ilustrados para os primeiros passos', 'CC BY 4.0 (20) e CC BY-NC 3.0 (2), impressas em cada PDF', 'PDF hospedado aqui'),
            ('MeloLingua', 'https://melolingua.com/learn-french', 'Histórias graduadas A1–B2 (níveis oficiais), com áudio na página oficial', 'Texto: CC BY-NC 4.0 (dataset). Áudio: só na página oficial', 'PDF aqui + link para ler e ouvir'),
            ('Project Gutenberg', 'https://www.gutenberg.org', 'Leitores escolares, contos, romances, teatro, poesia, filosofia, história e espiritualidade', 'Domínio público nos EUA', 'Baixar EPUB e Acessar (direto no Gutenberg)'),
            ('ELTeC-fra', 'https://github.com/COST-ELTeC/ELTeC-fra', '100 romances franceses de 1840 a 1920 (corpus acadêmico europeu)', 'CC BY 4.0', 'PDF hospedado aqui (99); 1 só com link'),
            ('Wikisource', 'https://fr.wikisource.org', 'Racine, Corneille, Molière, Pascal e a Chanson de Roland', 'CC BY-SA 4.0', 'PDF hospedado aqui, sob a mesma licença'),
            ('LibriVox', 'https://librivox.org', 'Audiolivros lidos por voluntários', 'Domínio público', 'Botão Audiolivro (link)')]
    tr = ''.join(f'<tr><td><a href="{u}">{n}</a><br><small>{sum(v for k, v in cnt.items() if k.startswith(n.split("-")[0].split(" ")[0])) or ""} {"obras" if n != "LibriVox" else ""}</small></td><td>{c}</td><td>{l}</td><td>{m}</td></tr>' for n, u, c, l, m in rows)
    body = f'''<div class="wrap narrow prose">
<h1>Sobre e direitos</h1>
<p class="lead">Existem ótimas leituras gratuitas em francês, mas espalhadas por muitos sites. A Bibliothèque Française reúne, organiza por nível e indica um caminho, do primeiro livrinho ilustrado aos clássicos.</p>
<p>É um projeto pessoal e sem fins lucrativos de um brasileiro que estuda francês, irmão da <a href="https://oaleffemanuel.github.io/biblioteca-italiana/">Biblioteca Italiana</a>. Não tem anúncios, não coleta dados e não é afiliado a nenhuma das fontes.</p>
<h2>As fontes</h2>
<div class="table-scroll"><table><thead><tr><th>Fonte</th><th>O que oferece</th><th>Licença</th><th>Como aparece aqui</th></tr></thead><tbody>{tr}</tbody></table></div>
<p>Também recomendamos, só como link: <a href="{INKLINGO}">Inklingo</a> (histórias A0–B2 com áudio; conteúdo protegido, leitura apenas no site deles).</p>
<h2 id="direitos">Política de direitos autorais</h2>
<ul><li><b>Acesso gratuito não é permissão.</b> Só hospedamos arquivos com licença verificada que permite redistribuição: ELTeC (CC BY 4.0), MeloLingua (CC BY-NC 4.0), African Storybook (CC BY 4.0 e CC BY-NC 3.0) e Wikisource (CC BY-SA 4.0). A evidência de cada obra está no <a href="donnees/manifesto.csv">manifesto de publicação</a>.</li>
<li><b>Domínio público nos EUA não é domínio público em todo lugar.</b> As obras do Project Gutenberg são acessadas diretamente no Gutenberg, que é a fonte oficial.</li>
<li><b>O original em domínio público não cobre traduções e edições modernas.</b> Por isso a coleção pessoal que deu origem ao projeto (traduções e edições de terceiros) não é publicada aqui.</li>
<li><b>Texto e áudio têm direitos separados.</b> Nenhum áudio é hospedado: os audiolivros são do LibriVox (domínio público) e os áudios MeloLingua ficam na página oficial.</li>
<li><b>Atribuição.</b> Cada PDF hospedado traz na primeira (ou última) página a fonte, a licença e os créditos; a lista completa está em <a href="livres/CREDITOS.md">livres/CREDITOS.md</a>. Os PDFs do ELTeC, do MeloLingua e do Wikisource foram gerados a partir dos arquivos oficiais, sem alterar o texto.</li></ul>
<p class="note">É titular de direitos e quer corrigir ou retirar algo? Abra uma <a href="{REPO}/issues">issue no GitHub</a>; a correção é feita com prioridade.</p>
<h2>Licença deste projeto</h2>
<ul><li><b>Textos do site, guia, planilha e dados do catálogo:</b> CC BY 4.0.</li>
<li><b>Código:</b> MIT.</li>
<li><b>PDFs em livres/:</b> a licença de cada fonte (CC BY 4.0, CC BY-NC 4.0, CC BY-NC 3.0 ou CC BY-SA 4.0), indicada no manifesto.</li></ul>
</div>'''
    return page('sobre.html', 'Sobre o projeto e direitos autorais', 'Fontes, licenças e a política de direitos autorais da Bibliothèque Française.', body)

def p_404():
    return page('404.html', 'Página não encontrada', 'Página não encontrada.',
                '<div class="wrap narrow"><h1>Page introuvable</h1><p>Essa página não existe. Volte ao <a href="./">início</a> ou ao <a href="catalogo.html">catálogo</a>.</p></div>')

# ───────────────────────── PDF do guia ─────────────────────────
PDF_CSS = '''
@page { size: A4; margin: 18mm 17mm 20mm; @bottom-center { content: "Bibliothèque Française · Guia de leitura · " counter(page); font: 8.5pt Helvetica, Arial, sans-serif; color: #7a8090; } }
@page :first { margin: 0; @bottom-center { content: none; } }
body { font: 10.4pt/1.5 Georgia, "Iowan Old Style", serif; color: #14213d; }
a { color: #b3263a; text-decoration: none; }
h1, h2, h3 { font-family: Georgia, serif; line-height: 1.2; }
h2 { font-size: 17pt; margin: 18pt 0 6pt; border-top: 1px solid #d8dae2; padding-top: 10pt; break-after: avoid; }
h3 { font-size: 12pt; margin: 12pt 0 4pt; break-after: avoid; }
p, li { orphans: 3; widows: 3; }
.lead { font-size: 12pt; color: #4a5468; }
table { width: 100%; border-collapse: collapse; font: 8.6pt/1.35 Helvetica, Arial, sans-serif; margin: 4pt 0 10pt; }
th { text-align: left; font-size: 7.4pt; text-transform: uppercase; letter-spacing: .04em; color: #4a5468; border-bottom: 1.2px solid #14213d; padding: 4pt 5pt; }
td { border-bottom: .5px solid #d8dae2; padding: 3.5pt 5pt; vertical-align: top; }
td:nth-child(3), td:nth-child(4) { white-space: nowrap; }
tr { break-inside: avoid; }
.keep { break-inside: avoid; }
.two { display: flex; gap: 8pt; flex-wrap: wrap; }
.two > .box { flex: 1 1 45%; }
.box { border: .8px solid #d8dae2; border-radius: 4pt; padding: 8pt 10pt; margin: 5pt 0; background: #f7f7f4; font-size: 9.6pt; break-inside: avoid; }
.note { background: #f1f1ed; border-left: 3pt solid #1f4e8c; padding: 6pt 10pt; font-size: 9.6pt; break-inside: avoid; }
.ext { font-size: 9pt; color: #4a5468; }
.tag { font: bold 8pt Helvetica, Arial, sans-serif; color: #fff; background: var(--c); padding: 2pt 7pt; border-radius: 9pt; }
.tappa-head { display: flex; align-items: baseline; gap: 8pt; border-top: 3pt double var(--c); margin-top: 6pt; }
.tappa-head h2 { border: 0; color: var(--c); margin-top: 8pt; }
.tappa-sec { break-before: page; }
#trilhas { break-before: page; }
.c1 { --c: #1f4e8c; } .c2 { --c: #2f6eb0; } .c3 { --c: #56709e; } .c4 { --c: #7d4c78; } .c5 { --c: #b3263a; } .c6 { --c: #14213d; }
.cover { width: 210mm; height: 297mm; background: #f8f7f3; padding: 30mm 24mm; box-sizing: border-box; position: relative; page-break-after: always; }
.cover .frame { position: absolute; top: 12mm; left: 12mm; right: 12mm; bottom: 12mm; border: 1px solid #14213d; outline: 1px solid #b3263a; outline-offset: -3mm; }
.cover .k { font: italic 14pt Georgia, serif; color: #b3263a; }
.cover h1 { font-size: 42pt; margin: 6pt 0 4pt; }
.cover .sub { font-size: 15pt; color: #4a5468; max-width: 140mm; }
.cover .tri { display: flex; height: 3mm; width: 60mm; margin: 14mm 0 10mm; }
.cover .tri div { flex: 1; }
.cover ol { font: 11pt/1.8 Helvetica, Arial, sans-serif; padding-left: 5mm; }
.cover .foot { position: absolute; bottom: 22mm; left: 24mm; right: 24mm; font: 9pt/1.5 Helvetica, Arial, sans-serif; color: #4a5468; }
.cover img { width: 18mm; }
.table-scroll { overflow: visible; }
'''
def build_pdf():
    from weasyprint import HTML
    import base64
    logo = 'data:image/svg+xml;base64,' + base64.b64encode(open(os.path.join(DOCS, 'assets', 'favicon.svg'), 'rb').read()).decode()
    toc = ''.join(f'<li>Étape {t} · <i>{e(d["nome"])}</i> — {e(d["pt"])} ({e(d["nivel"])})</li>' for t, d in ETAPES.items())
    doc = f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Guia de leitura · Bibliothèque Française</title><style>{PDF_CSS}</style></head><body>
<section class="cover"><div class="frame"></div><img src="{logo}" alt=""><p class="k">Bibliothèque Française · lire pour apprendre</p>
<h1>Guia de leitura</h1><p class="sub">Do A1 à literatura francesa original, em seis etapas, com as obras gratuitas da biblioteca, método e dicas para brasileiros.</p>
<div class="tri"><div style="background:#1f4e8c"></div><div style="background:#e6e6e0"></div><div style="background:#b3263a"></div></div><ol>{toc}</ol>
<p class="foot">Versão de {TODAY} · {SITE}<br>Texto sob licença CC BY 4.0. As obras citadas seguem as licenças de suas fontes; os links levam ao arquivo no site ou à fonte oficial.</p></section>
{plan_body(True)}</body></html>'''
    HTML(string=doc, base_url=DOCS).write_pdf(os.path.join(DOCS, 'telechargements', PDF_NAME))

# ───────────────────────── planilha ─────────────────────────
def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.formatting.rule import DataBarRule, FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    INK, ACC = '14213D', 'B3263A'
    TC = {1: '1F4E8C', 2: '2F6EB0', 3: '56709E', 4: '7D4C78', 5: 'B3263A', 6: '14213D'}
    hf, hfill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor=INK)
    thin = Border(bottom=Side(style='thin', color='D8DAE2'))
    wb = Workbook()
    def header(ws, cols, widths):
        for i, (c, wd) in enumerate(zip(cols, widths), 1):
            x = ws.cell(row=1, column=i, value=c); x.font = hf; x.fill = hfill; x.alignment = Alignment(vertical='center', wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = wd
        ws.row_dimensions[1].height = 30; ws.freeze_panes = 'A2'
    ws = wb.active; ws.title = 'Comece aqui'; ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3; ws.column_dimensions['B'].width = 104
    lines = [('Bibliothèque Française · Diário de leitura', Font(size=22, bold=True, color=ACC)), ('lire pour apprendre', Font(size=12, italic=True, color='4A5468')), ('', None),
             ('Como usar, em quatro passos', Font(size=14, bold=True)),
             ('1. Na aba "Catálogo", encontre uma obra (use os filtros do cabeçalho: Etapa, Nível, Áudio...).', None),
             ('2. Clique no link da coluna "Abrir" (PDF, EPUB ou página da fonte). Ao começar, mude o Status para "Lendo"; ao terminar, para "Lido".', None),
             ('3. Registre as sessões na aba "Sessões" e as palavras novas em "Vocabulário" (no máximo 5 por texto).', None),
             ('4. Escreva seu résumé en trois phrases na aba "Résumés". O "Painel" mostra seu progresso automaticamente.', None), ('', None),
             ('As seis etapas', Font(size=14, bold=True))]
    lines += [(f'Étape {t} · {d["nome"]} ({d["pt"]}, {d["nivel"]}): {d["curta"]}', None) for t, d in ETAPES.items()]
    lines += [('', None), ('Sobre os níveis', Font(size=14, bold=True)),
              ('• "Nível (oficial)" só aparece quando a fonte declara o CEFR (histórias MeloLingua). "Nível estimado" é uma estimativa do projeto.', None),
              ('• Se uma página tem mais de 5 a 8 palavras desconhecidas que atrapalham a história, desça uma etapa.', None), ('', None),
              (f'Site, catálogo e guia de leitura: {SITE}', Font(color=ACC, underline='single')),
              ('Planilha sob licença CC BY 4.0. As obras seguem as licenças de suas fontes.', Font(size=9, color='6B7280'))]
    for r, (txt, f) in enumerate(lines, 2):
        c = ws.cell(row=r, column=2, value=txt); c.alignment = Alignment(wrap_text=True, vertical='top')
        if f: c.font = f
        if txt.startswith('Site,'): c.hyperlink = SITE
        if txt.startswith('Étape '): c.font = Font(color=TC[int(txt.split()[1])])
    cat = wb.create_sheet('Catálogo')
    cols = ['Nº', 'Etapa', 'Nível (oficial)', 'Nível estimado', 'Título', 'Autor', 'Forma', 'Fonte', 'Tempo (min)', 'Áudio', 'Abrir', 'Status', 'Início', 'Fim', 'Nota (1–5)', 'Résumé?', 'Comentários']
    header(cat, cols, [5, 7, 10, 10, 40, 26, 18, 18, 10, 12, 10, 12, 11, 11, 9, 9, 36])
    for r, w in enumerate(W, 2):
        href = (SITE + w['arquivo']) if w['arquivo'] else ((w.get('epub_urls') or [None])[0] if len(w.get('epub_urls') or []) == 1 else w['url_fonte'])
        lab = 'PDF ↗' if w['arquivo'] else ('EPUB ↗' if len(w.get('epub_urls') or []) == 1 else 'Acessar ↗')
        au = 'Audiolivro' if audio_link(w) else ('Na fonte' if w['audio'] else '')
        vals = [r - 1, w['etapa'], w['cefr_oficial'] or '', '' if w['cefr_oficial'] else ETAPES[w['etapa']]['nivel'], w['t'], autor_line(w),
                w['forma'], w['fonte'], w['minutos'], au, lab]
        for ci, v in enumerate(vals, 1): cat.cell(row=r, column=ci, value=v).border = thin
        lk = cat.cell(row=r, column=11); lk.hyperlink = href; lk.font = Font(color=ACC, underline='single')
        t = cat.cell(row=r, column=2); t.fill = PatternFill('solid', fgColor=TC[w['etapa']]); t.font = Font(bold=True, color='FFFFFF'); t.alignment = Alignment(horizontal='center')
        for ci in (13, 14): cat.cell(row=r, column=ci).number_format = 'DD/MM/YYYY'
        for ci in range(12, 18): cat.cell(row=r, column=ci).border = thin
    last = len(W) + 1
    for rng, f in [(f'L2:L{last}', '"Quero ler,Lendo,Lido,Abandonei"'), (f'O2:O{last}', '"1,2,3,4,5"'), (f'P2:P{last}', '"Sim,Não"')]:
        dv = DataValidation(type='list', formula1=f, allow_blank=True); cat.add_data_validation(dv); dv.add(rng)
    dvd = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True, error='Digite uma data (ex.: 15/03/2027).', errorTitle='Data')
    cat.add_data_validation(dvd); dvd.add(f'M2:N{last}')
    for st, col in [('Lido', 'DDE7F5'), ('Lendo', 'FBEFD2'), ('Quero ler', 'EEF0F5'), ('Abandonei', 'EEEEEE')]:
        cat.conditional_formatting.add(f'A2:Q{last}', FormulaRule(formula=[f'$L2="{st}"'], fill=PatternFill('solid', fgColor=col)))
    cat.auto_filter.ref = f'A1:Q{last}'
    obra_list = f"'Catálogo'!$E$2:$E${last}"
    ses = wb.create_sheet('Sessões')
    header(ses, ['Data', 'Obra', 'Minutos', 'Páginas ou capítulos', 'Com áudio?', 'Observações'], [12, 42, 10, 18, 11, 50])
    d = DataValidation(type='list', formula1=f'={obra_list}', allow_blank=True, showErrorMessage=False); ses.add_data_validation(d); d.add('B2:B1000')
    d = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); ses.add_data_validation(d); d.add('E2:E1000')
    for r in range(2, 1001): ses.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
    voc = wb.create_sheet('Vocabulário')
    header(voc, ['Data', 'Mot', 'Tradução', 'Frase do texto', 'Obra', 'Revisar (1 dia)', 'Revisar (3 dias)', 'Revisar (7 dias)', 'Aprendida?'], [12, 18, 20, 50, 34, 13, 13, 13, 11])
    d = DataValidation(type='list', formula1=f'={obra_list}', allow_blank=True, showErrorMessage=False); voc.add_data_validation(d); d.add('E2:E2000')
    d = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=True); voc.add_data_validation(d); d.add('I2:I2000')
    for r in range(2, 501):
        voc.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
        for ci, n in ((6, 1), (7, 3), (8, 7)):
            c = voc.cell(row=r, column=ci, value=f'=IF(A{r}="","",A{r}+{n})'); c.number_format = 'DD/MM/YYYY'; c.font = Font(color='6B7280')
    voc.conditional_formatting.add('F2:H500', FormulaRule(formula=['AND(F2<>"",F2<=TODAY(),$I2<>"Sim")'], fill=PatternFill('solid', fgColor='FBEFD2')))
    res = wb.create_sheet('Résumés')
    header(res, ['Data', 'Obra', '1. Qui ? (quem e onde)', '2. Que se passe-t-il ? (o que acontece)', '3. Comment ça finit ? (fim ou opinião)'], [12, 34, 40, 48, 48])
    d = DataValidation(type='list', formula1=f'={obra_list}', allow_blank=True, showErrorMessage=False); res.add_data_validation(d); d.add('B2:B1000')
    for r in range(2, 301):
        res.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
        for ci in (3, 4, 5): res.cell(row=r, column=ci).alignment = Alignment(wrap_text=True, vertical='top')
    pn = wb.create_sheet('Painel', 1); pn.sheet_view.showGridLines = False
    for col, wd in zip('ABCDEFGH', [9, 30, 9, 9, 9, 11, 12, 4]): pn.column_dimensions[col].width = wd
    pn['A1'] = 'Seu progresso'; pn['A1'].font = Font(size=20, bold=True, color=ACC)
    pn['A2'] = 'Atualiza sozinho conforme você preenche as outras abas.'; pn['A2'].font = Font(italic=True, color='6B7280')
    kpis = [('Obras lidas', '=COUNTIF(Catálogo!$L:$L,"Lido")'), ('Lendo agora', '=COUNTIF(Catálogo!$L:$L,"Lendo")'),
            ('Lidas neste mês', '=COUNTIFS(Catálogo!$L:$L,"Lido",Catálogo!$N:$N,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1))'),
            ('Minutos de leitura (sessões)', '=SUM(Sessões!$C:$C)'),
            ('Dias com leitura (sessões)', '=SUMPRODUCT((Sessões!$A$2:$A$1000<>"")/COUNTIF(Sessões!$A$2:$A$1000,Sessões!$A$2:$A$1000&""))'),
            ('Palavras anotadas', '=COUNTA(Vocabulário!$B$2:$B$2000)'), ('Palavras aprendidas', '=COUNTIF(Vocabulário!$I:$I,"Sim")'),
            ('Résumés escritos', '=COUNTA(Résumés!$B$2:$B$1000)')]
    for i, (lab, f) in enumerate(kpis):
        r = 4 + i; pn.cell(row=r, column=1, value=lab).font = Font(color='4A5468'); pn.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = pn.cell(row=r, column=3, value=f); c.font = Font(size=13, bold=True); c.alignment = Alignment(horizontal='right')
    r0 = 4 + len(kpis) + 1
    for i, h in enumerate(['Etapa', 'Nome', 'Obras', 'Lidas', 'Lendo', 'Quero ler', '% lidas'], 1):
        c = pn.cell(row=r0, column=i, value=h); c.font = hf; c.fill = hfill
    for k, (t, dd) in enumerate(ETAPES.items(), 1):
        r = r0 + k
        a = pn.cell(row=r, column=1, value=t); a.fill = PatternFill('solid', fgColor=TC[t]); a.font = Font(bold=True, color='FFFFFF'); a.alignment = Alignment(horizontal='center')
        pn.cell(row=r, column=2, value=f'{dd["nome"]} ({dd["nivel"]})')
        pn.cell(row=r, column=3, value=f'=COUNTIF(Catálogo!$B:$B,A{r})')
        pn.cell(row=r, column=4, value=f'=COUNTIFS(Catálogo!$B:$B,A{r},Catálogo!$L:$L,"Lido")')
        pn.cell(row=r, column=5, value=f'=COUNTIFS(Catálogo!$B:$B,A{r},Catálogo!$L:$L,"Lendo")')
        pn.cell(row=r, column=6, value=f'=COUNTIFS(Catálogo!$B:$B,A{r},Catálogo!$L:$L,"Quero ler")')
        c = pn.cell(row=r, column=7, value=f'=IF(C{r}=0,0,D{r}/C{r})'); c.number_format = '0%'
        for ci in range(1, 8): pn.cell(row=r, column=ci).border = thin
    rt = r0 + len(ETAPES) + 1
    pn.cell(row=rt, column=2, value='Total').font = Font(bold=True)
    for ci, col in ((3, 'C'), (4, 'D'), (5, 'E'), (6, 'F')):
        pn.cell(row=rt, column=ci, value=f'=SUM({col}{r0+1}:{col}{rt-1})').font = Font(bold=True)
    c = pn.cell(row=rt, column=7, value=f'=IF(C{rt}=0,0,D{rt}/C{rt})'); c.number_format = '0%'; c.font = Font(bold=True)
    pn.conditional_formatting.add(f'G{r0+1}:G{rt}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='1F4E8C'))
    wb.properties.creator = 'Bibliothèque Française'
    wb.save(os.path.join(DOCS, 'telechargements', XLSX_NAME))

# ───────────────────────── dados abertos, créditos, sitemap ─────────────────────────
def build_data():
    os.makedirs(os.path.join(DOCS, 'donnees'), exist_ok=True)
    pub = [{k: v for k, v in w.items() if k not in ('i',)} for w in W]
    json.dump(pub, open(os.path.join(DOCS, 'donnees', 'catalogo.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    cols = ['id', 'etapa', 'cefr_oficial', 'dificuldade', 'nivel_fonte', 't', 'autor', 'tradutor', 'forma', 'fonte', 'minutos', 'paginas', 'arquivo', 'url_fonte', 'licenca', 'atribuicao']
    with open(os.path.join(DOCS, 'donnees', 'catalogo.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        wr = csv.writer(f); wr.writerow([c if c != 't' else 'titulo' for c in cols] + ['audio_url'])
        for w in W: wr.writerow([(SITE + w[c]) if c == 'arquivo' and w[c] else w[c] for c in cols] + [(w['audio'] or {}).get('url')])
    import shutil
    shutil.copy2(os.path.join(ROOT, 'data', 'manifesto.csv'), os.path.join(DOCS, 'donnees', 'manifesto.csv'))
    hosted = [w for w in W if w['arquivo']]
    rows = '\n'.join(f"| [{w['t']}]({os.path.basename(w['arquivo'])}) | {w['licenca']} | {w['atribuicao']} |" for w in sorted(hosted, key=lambda w: (w['fonte'], w['t'])))
    open(os.path.join(DOCS, 'livres', 'CREDITOS.md'), 'w', encoding='utf-8').write(f'''# Créditos e licenças dos PDFs hospedados

Todos os {len(hosted)} PDFs desta pasta têm licença que permite redistribuição. Cada PDF também traz a fonte e a licença na primeira página (ELTeC, MeloLingua, Wikisource) ou na última (African Storybook).

- **ELTeC-fra**: European Literary Text Collection, COST Action CA16204 "Distant Reading for European Literary History", DOI 10.5281/zenodo.3462535 — CC BY 4.0 — https://creativecommons.org/licenses/by/4.0/
- **MeloLingua**: CEFR-Graded Multilingual Stories Dataset 1.0.0 — CC BY-NC 4.0 — https://creativecommons.org/licenses/by-nc/4.0/ (marca MeloLingua não licenciada)
- **African Storybook**: www.africanstorybook.org — CC BY 4.0 ou CC BY-NC 3.0, conforme cada livro
- **Wikisource (fr)**: transcrições sob CC BY-SA 4.0; os PDFs derivados seguem a mesma licença — https://creativecommons.org/licenses/by-sa/4.0/

Mudanças: os PDFs do ELTeC, do MeloLingua e do Wikisource foram gerados a partir dos arquivos oficiais (TEI, Markdown, HTML), sem alterar o texto. Os PDFs do African Storybook são os oficiais.

| Obra | Licença | Atribuição |
|---|---|---|
{rows}
''')
    urls = ['', 'catalogo.html', 'plano.html', 'guia.html', 'sobre.html']
    open(os.path.join(DOCS, 'sitemap.xml'), 'w').write('<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
        + ''.join(f'<url><loc>{SITE}{u}</loc><lastmod>{TODAY}</lastmod></url>' for u in urls) + '</urlset>\n')
    open(os.path.join(DOCS, 'robots.txt'), 'w').write(f'User-agent: *\nAllow: /\nSitemap: {SITE}sitemap.xml\n')
    open(os.path.join(DOCS, '.nojekyll'), 'w').write('')

if __name__ == '__main__':
    os.makedirs(os.path.join(DOCS, 'telechargements'), exist_ok=True)
    for fname, fn in [('index.html', p_index), ('catalogo.html', p_catalogo), ('plano.html', p_plano), ('guia.html', p_guia), ('sobre.html', p_sobre), ('404.html', p_404)]:
        open(os.path.join(DOCS, fname), 'w', encoding='utf-8').write(fn())
    build_data(); build_xlsx()
    if '--sem-pdf' not in sys.argv: build_pdf()
    print(f'ok: {N} obras · {NA} com áudio · {NP} PDFs hospedados · SITE={SITE}')
